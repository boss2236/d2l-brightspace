# SPDX-License-Identifier: AGPL-3.0-or-later
"""Plain text out of course files, so search and the AI can read lecture notes. Unknown types return ''."""
import html
import logging
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

logging.getLogger("pypdf").setLevel(logging.ERROR)     # font-encoding chatter on every maths PDF

# what gets downloaded at all; videos and archives are skipped
TEXT_TYPES = {"pdf", "docx", "pptx", "html", "htm", "txt", "md", "csv"}
MAX_BYTES = 40 * 1024 * 1024


def _tidy(t: str) -> str:
    """pypdf writes some PDFs one word per line (with ' ' lines between); join those pages back into prose."""
    t = re.sub(r"(\s?\.){5,}", " … ", t)                   # table-of-contents dot leaders
    lines = [re.sub(r"[ \t]{2,}", " ", l).strip() for l in t.split("\n")]
    words = [l for l in lines if l]
    if len(words) > 8 and sum(len(l.split()) for l in words) / len(words) < 1.3:
        return " ".join(words)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _html_text(raw: str) -> str:
    raw = raw.replace("\r", "").lstrip("\ufeff")
    raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
    # block elements end a line even when the HTML source doesn't
    raw = re.sub(r"(?i)<br\s*/?>|</(p|div|h[1-6]|li|tr|table|section|article|blockquote|pre)>", "\n", raw)
    text = html.unescape(re.sub(r"<[^>]+>", " ", raw)).replace("\xa0", " ").lstrip("\ufeff")
    return re.sub(r"[ \t]*\n\s*", "\n", re.sub(r"[ \t]{2,}", " ", text)).strip()


def ocr_available() -> bool:
    return os.environ.get("D2L_OCR", "1") != "0" and bool(shutil.which("tesseract") and shutil.which("pdftoppm"))


def _ocr(path: Path, page_numbers: list[int]) -> dict[int, str]:
    """Text of scanned PDF pages via poppler's pdftoppm + tesseract, both run as plain subprocesses (no shell) with
    time limits. Off when D2L_OCR=0 or the tools aren't installed; D2L_OCR_LANG picks languages (default eng),
    D2L_OCR_MAX_PAGES caps work per file (default 40)."""
    if not ocr_available():
        return {}
    lang = re.sub(r"[^A-Za-z_+]", "", os.environ.get("D2L_OCR_LANG", "eng")) or "eng"
    out: dict[int, str] = {}
    with tempfile.TemporaryDirectory() as tmp:
        for n in page_numbers[: int(os.environ.get("D2L_OCR_MAX_PAGES", "40"))]:
            img = Path(tmp) / f"p{n}"
            try:
                subprocess.run(["pdftoppm", "-r", "200", "-gray", "-png", "-singlefile", "-f", str(n), "-l", str(n),
                                str(path), str(img)], check=True, capture_output=True, timeout=60)
                r = subprocess.run(["tesseract", f"{img}.png", "-", "-l", lang], check=True, capture_output=True,
                                   timeout=120)
            except (subprocess.SubprocessError, OSError):
                continue
            out[n] = _tidy(r.stdout.decode("utf-8", "ignore"))
    return out


def text_of(path: Path) -> str:
    ext = path.suffix.lower().lstrip(".")
    if ext not in ("docx", "pptx") and zipfile.is_zipfile(path):
        # Brightspace serves HTML topics as a zip of the page plus its assets
        with zipfile.ZipFile(path) as z:
            pages = [n for n in z.namelist() if n.lower().endswith((".html", ".htm", ".txt"))]
            return "\n\n".join(_html_text(z.read(n).decode("utf-8", "ignore")) for n in pages).strip()
    try:
        if ext == "pdf":
            from pypdf import PdfReader
            pages = [_tidy(p.extract_text() or "") for p in PdfReader(path).pages]
            thin = [i for i, t in enumerate(pages) if len(t.split()) < 12]     # near-empty = a scanned page
            ocr = _ocr(path, [i + 1 for i in thin]) if thin else {}
            for i, t in ocr.items():
                if len(t.split()) > len(pages[i - 1].split()):
                    pages[i - 1] = t
            text = "\n\n".join(f"[page {i}]\n{t}" for i, t in enumerate(pages, 1)).strip()
            words = len(re.sub(r"\[page \d+\]", "", text).split())
            if ocr:
                text = "[some pages were scanned images; their text was recognised by OCR and may contain errors]\n\n" + text
            if words < 25 * max(len(pages), 1):
                text = ("[mostly images or handwriting: only a little text could be extracted — open the file to read it]\n\n"
                        + text)
            return text
        if ext == "docx":
            from docx import Document
            return "\n".join(p.text for p in Document(path).paragraphs if p.text.strip())
        if ext == "pptx":
            from pptx import Presentation
            slides = []
            for i, s in enumerate(Presentation(path).slides, 1):
                bits = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
                slides.append(f"[slide {i}]\n" + "\n".join(bits))
            return "\n\n".join(slides)
        if ext in ("html", "htm"):
            return _html_text(path.read_text(errors="ignore"))
        if ext in ("txt", "md", "csv"):
            return path.read_text(errors="ignore")
    except Exception as e:                      # a broken file shouldn't stop the sync
        return f"[could not extract text: {e.__class__.__name__}]"
    return ""


def reindex(db, root: Path) -> int:
    """Re-extract text from files already downloaded (after improving the extractor); no network."""
    n = 0
    for r in db.execute("SELECT id, file_path FROM content WHERE file_status = 'ok'").fetchall():
        db.execute("UPDATE content SET text = ? WHERE id = ?", (text_of(root / r[1]), r[0]))
        n += 1
    return n
